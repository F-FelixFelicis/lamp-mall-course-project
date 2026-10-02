from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\软件工程\灯具商城项目")
TOOLS = ROOT / "tools"
sys.path.insert(0, str(TOOLS))

from build_srs import (  # noqa: E402
    NAVY,
    PALE_BLUE,
    TEXT_GRAY,
    add_body,
    add_bullet,
    add_heading,
    add_page_field,
    add_table,
    configure_sections,
    configure_styles,
    remove_paragraph_borders,
    set_cell_borders,
    set_cell_margins,
    set_run_font,
    shade_cell,
)


OUTPUT = ROOT / "文档" / "灯具商城系统设计说明书_v0.1.docx"


def add_footer(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("灯具商城系统设计说明书 v0.1    第 ")
    set_run_font(run, east_asia="宋体", size=9, color=TEXT_GRAY)
    add_page_field(paragraph)
    run = paragraph.add_run(" 页")
    set_run_font(run, east_asia="宋体", size=9, color=TEXT_GRAY)


def add_center_line(doc, text, size=11, bold=False, color=None, before=0, after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_run_font(r, east_asia="微软雅黑" if bold else "宋体", size=size, bold=bold, color=color)
    return p


def add_note(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.9)
    set_cell_borders(cell, color="9FBAD0", size="8")
    set_cell_margins(cell, top=110, start=140, bottom=110, end=140)
    shade_cell(cell, PALE_BLUE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(title + "：")
    set_run_font(r, east_asia="微软雅黑", size=9.6, bold=True, color=RGBColor(32, 56, 100))
    r = p.add_run(text)
    set_run_font(r, east_asia="宋体", size=9.6)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_borders(cell, color="BFBFBF", size="6")
    set_cell_margins(cell, top=110, start=140, bottom=110, end=140)
    shade_cell(cell, "F4F6F8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    for idx, line in enumerate(lines):
        r = p.add_run(line + ("\n" if idx < len(lines) - 1 else ""))
        set_run_font(r, east_asia="等线", latin="Consolas", size=8.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


_MAJOR_BREAK_COUNT = 0


def add_major_break(doc):
    """Reserve hard breaks for the cover and front matter only.

    Body chapters flow naturally so a table continuation does not leave a
    mostly empty page before the next chapter.
    """
    global _MAJOR_BREAK_COUNT
    _MAJOR_BREAK_COUNT += 1
    if _MAJOR_BREAK_COUNT <= 2:
        doc.add_page_break()


def build():
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    add_footer(doc.sections[0])

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(76)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("灯具商城系统设计说明书")
    set_run_font(r, east_asia="微软雅黑", size=27, bold=True)
    p.style = doc.styles["Title"]
    remove_paragraph_borders(p)

    add_center_line(doc, "三端商城与以图搜图一体化设计", size=15, color=TEXT_GRAY, before=8)
    add_center_line(doc, "SYSTEM DESIGN SPECIFICATION", size=10, color=TEXT_GRAY, before=5)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in (
        "文档版本  v0.1",
        "编制日期  2026年9月24日",
        "设计范围  客户端小程序 / 商家端小程序 / PC管理端 / 图像检索",
    ):
        r = p.add_run(line + "\n")
        set_run_font(r, east_asia="宋体", size=11)

    add_center_line(
        doc,
        "本版作为数据库、接口、代码骨架和测试用例的共同技术基线",
        size=10.5,
        color=TEXT_GRAY,
        before=60,
    )

    add_major_break(doc)
    add_heading(doc, "文档控制", 1)
    add_table(
        doc,
        ["版本", "日期", "状态", "变更说明"],
        [["v0.1", "2026-09-24", "设计基线", "依据课程要求、灯具商城附件和需求规格说明书形成第一版可实现设计。"]],
        [0.75, 1.15, 1.0, 4.2],
        {0, 1, 2},
    )
    add_heading(doc, "设计输入与配套产物", 2)
    add_table(
        doc,
        ["产物", "位置", "用途"],
        [
            ["需求规格说明书", "文档/灯具商城需求规格说明书_v0.1.docx", "定义范围、角色、业务规则与验收标准。"],
            ["数据库脚本", "database/schema.sql", "MySQL 8.0 初始化表、索引、约束与基础角色。"],
            ["接口契约", "api/openapi-v0.1.yaml", "三端共同使用的 OpenAPI 3.1 核心接口契约。"],
            ["本说明书", "文档/灯具商城系统设计说明书_v0.1.docx", "统一架构、模块、事务、检索、测试和部署方案。"],
        ],
        [1.2, 3.1, 2.8],
        {0},
    )
    add_heading(doc, "目录", 2)
    for item in (
        "1 设计目标与约束",
        "2 总体架构与技术选型",
        "3 模块与代码边界",
        "4 数据库设计",
        "5 核心业务与状态机",
        "6 API 与前后端协作",
        "7 以图搜图子系统",
        "8 安全、隐私与审计",
        "9 部署与运行配置",
        "10 测试与质量保证",
        "11 需求追踪与迭代计划",
    ):
        add_bullet(doc, item)

    add_major_break(doc)
    add_heading(doc, "1 设计目标与约束", 1)
    add_heading(doc, "1.1 设计目标", 2)
    add_body(doc, "本设计面向课程综合类项目的第一版实现，目标不是堆叠功能，而是以可运行、可测试、可讲解的方式跑通商家入驻、灯具商品与 SKU、商家报价、审核、检索、下单、履约和评价的业务闭环，并将以图搜图作为系统内部可替换、可评估的算法模块。")
    for text in (
        "业务闭环完整：三个角色均有明确入口、权限、状态和操作结果。",
        "数据一致：订单保存商品、规格、商家和成交价格快照，历史记录不依赖当前报价。",
        "算法可复现：图像特征、索引版本、实验指标和线上结果可以追溯。",
        "适合团队协作：接口契约先行，数据库迁移受控，模块边界与负责人清晰。",
        "适合课程答辩：关键流程可演示，异常路径可说明，代码能够现场定位和修改。",
    ):
        add_bullet(doc, text)
    add_heading(doc, "1.2 关键约束", 2)
    add_table(
        doc,
        ["约束", "设计响应"],
        [
            ["课程周期与团队规模有限", "第一版采用模块化单体，减少分布式部署和跨服务事务成本。"],
            ["需要三类用户端", "客户与商家使用 uni-app 小程序，管理员使用 Vue 3 PC 端，共用 FastAPI 后端。"],
            ["包含算法任务", "图像检索保持独立模块接口，离线建索引、在线查 Top-K，业务层负责可售过滤。"],
            ["附件包含跨领域原型示例", "领域模型统一采用灯具概念，增加功率、色温、光通量、安装方式和智能协议等属性。"],
            ["第一版不接真实支付", "订单从创建进入待确认，不伪造支付回调；在线支付留作后续扩展。"],
        ],
        [1.65, 5.45],
        {0},
    )
    add_note(doc, "核心决策", "先完成可运行的模块化单体，再根据压力测试或团队分工需要拆服务；以图搜图虽然逻辑独立，但第一版可与后端同机部署。")

    add_major_break(doc)
    add_heading(doc, "2 总体架构与技术选型", 1)
    add_heading(doc, "2.1 逻辑架构", 2)
    add_table(
        doc,
        ["接入层", "应用层", "领域与算法层", "基础设施层"],
        [
            ["客户小程序\n商家小程序\nPC 管理端", "REST API\n认证鉴权\n参数校验\n异常映射", "账号权限\n商品报价\n订单履约\n审核审计\n图像检索", "MySQL 8\nRedis（可选）\n对象存储\nFAISS 索引\n日志"],
        ],
        [1.65, 1.75, 1.9, 1.8],
        {0, 1, 2, 3},
    )
    add_body(doc, "客户端不直接访问数据库、对象存储私有文件或检索索引。所有业务请求进入统一 API；API 完成身份识别、权限校验和参数验证后调用应用服务。应用服务编排领域规则与事务，基础设施适配数据库、缓存、文件和向量索引。")
    add_heading(doc, "2.2 技术栈", 2)
    add_table(
        doc,
        ["层次", "建议技术", "采用理由"],
        [
            ["客户/商家端", "uni-app + Vue 3 + TypeScript", "一套技术栈适配微信小程序，组件与请求层可共享。"],
            ["管理端", "Vue 3 + TypeScript + Vite + Pinia", "适合表格、审核、权限菜单和快速开发。"],
            ["后端", "Python 3.12 + FastAPI + Pydantic", "接口文档友好，便于与图像检索共用 Python 生态。"],
            ["数据访问", "SQLAlchemy 2 + Alembic", "显式事务、模型映射与可追踪数据库迁移。"],
            ["主数据库", "MySQL 8.0", "满足结构化交易数据、约束、索引与常用部署环境。"],
            ["缓存", "Redis 7（可选）", "验证码、短时令牌、幂等结果和热点缓存；不可作为订单事实来源。"],
            ["文件", "MinIO 或本地兼容对象存储", "业务表仅保存文件标识和元数据，便于后续迁移。"],
            ["图像检索", "PyTorch/ONNX Runtime + FAISS", "支持特征提取、余弦相似度检索和离线评估。"],
            ["测试", "pytest + httpx + Vitest", "覆盖领域规则、接口集成与前端工具函数。"],
        ],
        [1.15, 2.35, 3.8],
        {0},
    )
    add_heading(doc, "2.3 运行时请求链", 2)
    add_code(doc, [
        "小程序 / 管理端 → HTTPS /api/v1 → FastAPI 路由",
        "→ 身份与 RBAC → 应用服务 → SQLAlchemy 事务 → MySQL",
        "                         ↘ 文件服务 → MinIO",
        "                         ↘ 图像检索端口 → 特征模型 + FAISS",
    ])
    add_note(doc, "边界规则", "路由层不写业务状态，仓储层不决定审核结果，图像检索模块不决定商品是否可售；这些规则分别属于应用层、领域层和商城业务层。")

    add_major_break(doc)
    add_heading(doc, "3 模块与代码边界", 1)
    add_heading(doc, "3.1 后端模块", 2)
    add_table(
        doc,
        ["模块", "职责", "主要依赖"],
        [
            ["identity", "注册登录、令牌、用户、角色、密码与当前用户。", "users / roles / user_roles"],
            ["merchant", "商家申请、认证资料、审核记录和商家数据范围。", "identity / files"],
            ["catalog", "分类、SPU、SKU、图片、上下架与商品审核。", "merchant / files"],
            ["pricing", "商家报价、审核、生效区间、最低有效报价。", "catalog / merchant"],
            ["search", "文字筛选、图片预处理、特征检索与结果聚合。", "catalog / pricing / retrieval"],
            ["order", "下单、价格快照、状态流转、库存校验与履约。", "identity / pricing / address"],
            ["review", "订单完成后的文字与图片评价。", "order / files"],
            ["notification", "审核和订单事件形成站内通知。", "各应用服务发布的领域事件"],
            ["audit", "操作日志、审计查询和关键变更留痕。", "当前用户与请求上下文"],
            ["files", "上传校验、存储键、访问级别和签名访问。", "对象存储"],
        ],
        [1.15, 3.8, 2.35],
        {0},
    )
    add_heading(doc, "3.2 推荐代码目录", 2)
    add_code(doc, [
        "apps/",
        "  api/app/{core,identity,merchant,catalog,pricing,search,order,review}/",
        "  admin-web/src/{api,stores,views,components}/",
        "  customer-mini/src/{api,stores,pages,components}/",
        "  merchant-mini/src/{api,stores,pages,components}/",
        "packages/retrieval/{extractor,index,service,evaluation}/",
        "database/{schema.sql,migrations,seeds}/   api/openapi-v0.1.yaml",
        "tests/{unit,integration,e2e,retrieval}/   docs/",
    ])
    add_heading(doc, "3.3 依赖方向", 2)
    for text in (
        "路由层依赖应用服务；应用服务依赖领域对象与抽象仓储；基础设施实现仓储接口。",
        "订单模块可读取报价，但报价模块不得反向依赖订单；通过事件通知报价或运营统计。",
        "检索模块返回 SKU 编号与相似度，商城应用层再查询商品状态、有效报价和展示信息。",
        "前端统一从 OpenAPI 生成或手写类型化请求层，页面不得散落拼接接口地址。",
        "跨模块写操作由一个应用服务明确发起，避免同一请求中多处隐式提交事务。",
    ):
        add_bullet(doc, text)
    add_heading(doc, "3.4 三端页面边界", 2)
    add_table(
        doc,
        ["端", "第一版页面"],
        [
            ["客户端小程序", "登录、首页、分类/搜索、以图搜图、商品详情/报价、收藏、下单、订单列表/详情、评价、个人资料/地址。"],
            ["商家端小程序", "登录、认证申请、商品列表/编辑、SKU 与图片、报价、订单列表/详情、确认订单、发货凭证、消息。"],
            ["PC 管理端", "登录、仪表盘、用户/角色、商家审核、商品审核、报价审核、订单、评价、分类、操作日志。"],
        ],
        [1.4, 5.7],
        {0},
    )

    add_major_break(doc)
    add_heading(doc, "4 数据库设计", 1)
    add_heading(doc, "4.1 数据分组", 2)
    add_table(
        doc,
        ["数据域", "数据表", "说明"],
        [
            ["身份与文件", "users, roles, user_roles, uploaded_files", "账号、RBAC 和受控文件元数据。"],
            ["商家", "merchant_profiles, merchant_audits", "认证资料与不可覆盖的审核历史。"],
            ["目录", "categories, products, product_skus, product_images, product_audits", "灯具分类、SPU/SKU、图片和审核。"],
            ["报价与检索", "merchant_offers, offer_audits, image_embeddings, favorites", "多商家报价、向量元数据与收藏。"],
            ["交易与履约", "orders, order_items, shipment_records", "订单头、成交快照和发货记录。"],
            ["评价与运营", "reviews, review_images, notifications, operation_logs", "评价、消息和关键操作审计。"],
        ],
        [1.25, 3.35, 2.5],
        {0},
    )
    add_heading(doc, "4.2 核心关系", 2)
    add_code(doc, [
        "users 1─1 merchant_profiles 1─N merchant_offers N─1 product_skus",
        "categories 1─N products 1─N product_skus 1─N product_images",
        "users(customer) 1─N orders 1─N order_items N─1 merchant_offers",
        "order_items 1─0..1 reviews 1─N review_images",
        "product_images 1─N image_embeddings（按模型与索引版本区分）",
    ])
    add_heading(doc, "4.3 数据建模规则", 2)
    for text in (
        "所有金额使用 DECIMAL(12,2)，API 使用十进制字符串或受控数值，不使用二进制浮点计算成交金额。",
        "商品使用 SPU/SKU 两层模型；可选择的颜色、尺寸、功率、色温等落在 SKU，描述性公共信息落在 SPU。",
        "订单明细保存商品名、SKU 编码、规格、商家名、单价和报价编号快照，历史订单不随商品改名或报价失效而变化。",
        "用户、商品、文件等主数据使用 deleted_at 软删除；审核、订单、发货和操作日志不得物理删除。",
        "审核记录采用追加模式；业务对象保存当前状态，审核表保存每次决定及理由。",
        "uploaded_files 保存 SHA-256、MIME、大小和访问级别；私有文件通过授权接口或短时签名 URL 访问。",
    ):
        add_bullet(doc, text)
    add_heading(doc, "4.4 索引与约束", 2)
    add_table(
        doc,
        ["场景", "索引/约束策略"],
        [
            ["登录定位", "users.username 与 users.phone 唯一索引；至少一个标识不为空。"],
            ["商品列表", "products(category_id, status)、product_skus(product_id, status)。"],
            ["有效报价", "merchant_offers(sku_id, status, price) 与商家维度索引；价格和库存非负 CHECK。"],
            ["订单查询", "orders(customer_user_id, status, created_at)、订单号唯一索引。"],
            ["防重复评价", "reviews.order_item_id 唯一约束。"],
            ["向量版本", "image_embeddings 对图片、模型版本、索引版本建立唯一组合约束。"],
        ],
        [1.5, 5.6],
        {0},
    )
    add_note(doc, "实现提示", "schema.sql 是基线结构；进入开发后由 Alembic 管理增量迁移。禁止开发成员直接手工修改共享数据库后不提交迁移脚本。")

    add_major_break(doc)
    add_heading(doc, "5 核心业务与状态机", 1)
    add_heading(doc, "5.1 状态定义", 2)
    add_table(
        doc,
        ["对象", "状态", "允许变化"],
        [
            ["商家认证", "PENDING / APPROVED / REJECTED", "提交进入 PENDING；管理员决定 APPROVED 或 REJECTED；驳回后修改再提交。"],
            ["商品", "DRAFT / PENDING / APPROVED / ON_SALE / OFF_SALE / REJECTED", "草稿提交审核；通过后上架；下架可再次编辑提交。"],
            ["报价", "DRAFT / PENDING / ACTIVE / REJECTED / EXPIRED", "待审核通过后生效；新报价生效时旧报价失效。"],
            ["订单", "PENDING_CONFIRM / PENDING_SHIPMENT / SHIPPED / COMPLETED / CANCELLED", "商家确认、发货、客户确认完成；取消只允许在未发货前。"],
        ],
        [1.05, 2.4, 3.65],
        {0},
    )
    add_heading(doc, "5.2 创建订单事务", 2)
    add_code(doc, [
        "BEGIN",
        "1. 校验 Idempotency-Key，已成功则返回原订单",
        "2. 锁定并读取目标 ACTIVE 报价，校验商品/商家/库存/有效期",
        "3. 计算金额并写 orders；复制地址与收货信息快照",
        "4. 写 order_items；复制商品、SKU、商家与成交价格快照",
        "5. 扣减或预占库存；写操作日志与通知事件",
        "COMMIT（任一步失败则全部回滚）",
    ])
    add_heading(doc, "5.3 并发与幂等", 2)
    for text in (
        "创建订单要求 Idempotency-Key；同一用户、同一键、同一请求体只产生一个订单。",
        "库存更新使用条件更新或行锁，更新影响行数为零时返回库存不足，避免超卖。",
        "商品和报价编辑携带 version；更新时比较版本，冲突返回 409 并提示刷新。",
        "审核接口只接受待审核状态；重复提交相同审核决定返回当前结果，不重复生成业务副作用。",
        "文件上传先完成校验和存储，再在事务中关联业务对象；失败的临时文件由定时任务清理。",
    ):
        add_bullet(doc, text)
    add_heading(doc, "5.4 关键业务规则", 2)
    add_table(
        doc,
        ["编号", "规则"],
        [
            ["BR-01", "只有认证通过的商家可以提交商品、报价和处理本店订单。"],
            ["BR-02", "客户端只展示已上架商品、有效 SKU 和当前有效报价。"],
            ["BR-03", "同一商家同一 SKU 同时最多一条 ACTIVE 报价。"],
            ["BR-04", "审核驳回必须填写原因；前端需展示最近一次驳回原因。"],
            ["BR-05", "订单成交金额以后端重算结果为准，不能信任客户端传入总价。"],
            ["BR-06", "只有已完成订单的对应明细可以评价，且每条明细仅一次有效评价。"],
        ],
        [1.0, 6.1],
        {0},
    )

    add_major_break(doc)
    add_heading(doc, "6 API 与前后端协作", 1)
    add_heading(doc, "6.1 通用约定", 2)
    add_table(
        doc,
        ["项目", "约定"],
        [
            ["基础路径", "/api/v1；资源名使用复数名词。"],
            ["认证", "Authorization: Bearer <access_token>；后端同时校验角色与数据归属。"],
            ["时间", "ISO 8601，服务端统一存 UTC，展示层转换为 Asia/Shanghai。"],
            ["分页", "page 从 1 开始，page_size 默认 20、最大 100；返回 items/total/page/page_size。"],
            ["金额", "以两位小数表达；后端使用 Decimal，响应中保持精度。"],
            ["错误", "统一 code/message/details/request_id；校验错误返回字段级 details。"],
            ["并发", "编辑资源使用 version；冲突返回 409。创建订单使用 Idempotency-Key。"],
            ["文件", "上传接口限制 MIME、后缀、文件大小和图片像素，私有文件不返回永久公开地址。"],
        ],
        [1.25, 5.85],
        {0},
    )
    add_heading(doc, "6.2 第一版核心接口", 2)
    add_table(
        doc,
        ["模块", "接口", "说明"],
        [
            ["认证", "POST /auth/register\nPOST /auth/login\nGET /auth/me", "注册、登录和当前用户。"],
            ["目录", "GET /products\nGET /products/{product_id}", "可售商品列表、SKU 与有效报价详情。"],
            ["检索", "POST /search/image", "上传图片并返回相似 SKU Top-K。"],
            ["商家", "PUT /merchant/application\nPOST /merchant/products\nPOST /merchant/skus/{sku_id}/offers", "认证申请、商品草稿与报价。"],
            ["订单", "POST/GET /orders\nPOST /merchant/orders/{order_id}/confirm\nPOST /merchant/orders/{order_id}/shipments", "下单、查询、确认与发货。"],
            ["评价", "POST /order-items/{order_item_id}/review", "对已完成订单明细评价。"],
            ["管理", "POST /admin/merchants/{id}/audit\nPOST /admin/products/{id}/audit\nPOST /admin/offers/{id}/audit", "三类审核入口。"],
        ],
        [0.9, 3.55, 2.65],
        {0},
    )
    add_heading(doc, "6.3 错误码", 2)
    add_table(
        doc,
        ["HTTP", "业务码示例", "前端处理"],
        [
            ["400", "BAD_REQUEST", "提示请求错误，不自动重试。"],
            ["401", "UNAUTHORIZED", "尝试刷新令牌或回到登录页。"],
            ["403", "FORBIDDEN / MERCHANT_NOT_APPROVED", "隐藏无权限动作并显示原因。"],
            ["404", "RESOURCE_NOT_FOUND", "显示不存在或已下架。"],
            ["409", "VERSION_CONFLICT / OFFER_CHANGED / OUT_OF_STOCK", "刷新最新数据，让用户确认。"],
            ["422", "VALIDATION_ERROR", "在对应字段展示校验信息。"],
            ["500", "INTERNAL_ERROR", "显示 request_id，便于后台查日志。"],
        ],
        [0.8, 2.75, 3.55],
        {0},
    )
    add_note(doc, "契约管理", "openapi-v0.1.yaml 是第一版接口真源。后端变更接口时先更新契约与测试，前端基于契约同步类型，禁止只在聊天中口头约定字段。")

    add_major_break(doc)
    add_heading(doc, "7 以图搜图子系统", 1)
    add_heading(doc, "7.1 功能边界", 2)
    add_body(doc, "以图搜图负责把客户上传的灯具图片转换为特征向量，在已入库的商品主图或 SKU 图中检索最相似候选，并返回 SKU 编号、相似度和索引版本。它不负责商品审核、报价计算或可售判断；业务层需要对候选结果进行状态和报价过滤。")
    add_heading(doc, "7.2 离线建库流程", 2)
    add_code(doc, [
        "已审核并上架的商品图片",
        "→ 下载与格式校验 → 去重 → 统一 RGB / 尺寸 / 归一化",
        "→ 特征提取模型 → L2 归一化向量",
        "→ 写 image_embeddings 元数据 → 构建 FAISS 索引",
        "→ 原子发布 index_version → 生成覆盖率与检索质量报告",
    ])
    add_heading(doc, "7.3 在线检索流程", 2)
    add_code(doc, [
        "上传图片 → 类型/大小/像素校验 → EXIF 方向纠正 → RGB 预处理",
        "→ 提取查询向量 → FAISS 余弦相似度 Top-N",
        "→ SKU 去重与聚合 → 过滤下架/无效报价 → 返回 Top-K",
    ])
    add_heading(doc, "7.4 第一版模型与索引", 2)
    add_table(
        doc,
        ["项目", "第一版方案", "可替换点"],
        [
            ["特征模型", "预训练视觉编码器，先冻结特征，建立可复现实验基线。", "后续用灯具数据微调或度量学习。"],
            ["向量", "L2 归一化，保存模型版本、维度、checksum。", "切换模型时新建索引版本，不覆盖旧向量。"],
            ["相似度", "内积实现余弦相似度。", "可比较欧氏距离或融合属性重排。"],
            ["索引", "数据量较小时 IndexFlatIP；规模增长再评估 HNSW/IVF。", "以召回率、延迟和内存共同决策。"],
            ["结果", "先取 Top-N，再按 SKU 聚合并进行业务过滤。", "增加分类约束、品牌/风格重排。"],
        ],
        [1.1, 3.55, 2.45],
        {0},
    )
    add_heading(doc, "7.5 数据集与评估", 2)
    add_table(
        doc,
        ["指标", "定义", "第一版用途"],
        [
            ["Recall@K", "同款或同 SKU 正样本是否出现在前 K。", "主要离线检索指标。"],
            ["mAP@K", "多个相关结果的排序质量。", "比较模型和重排策略。"],
            ["P95 延迟", "在线检索 95% 请求的耗时。", "确保演示和交互可接受。"],
            ["索引覆盖率", "可售 SKU 中至少有一张有效向量的比例。", "监控漏建索引。"],
            ["失败率", "图片校验、特征提取或索引查询失败占比。", "定位数据与运行问题。"],
        ],
        [1.2, 3.45, 2.45],
        {0},
    )
    for text in (
        "数据集按商品或 SKU 分组切分，避免同一商品不同角度同时落入训练集和测试集造成泄漏。",
        "保留查询图、正样本定义、模型版本、索引版本、随机种子与评估脚本参数。",
        "必须设置关键词/分类基线与视觉基线，不能只展示若干成功截图。",
        "无明显相似结果时返回低置信提示，不强行把不相关商品包装为准确匹配。",
    ):
        add_bullet(doc, text)

    add_major_break(doc)
    add_heading(doc, "8 安全、隐私与审计", 1)
    add_heading(doc, "8.1 认证与授权", 2)
    for text in (
        "密码使用 Argon2id 或 bcrypt 哈希，禁止明文或可逆加密保存。",
        "访问令牌短时有效；刷新令牌可撤销并与设备或会话记录关联。",
        "后端按角色和资源归属双重校验，例如商家只能处理 merchant_id 属于自己的订单。",
        "管理员账号启用强密码，关键审核操作记录操作者、对象、前后状态、IP 和 request_id。",
        "登录、验证码、上传和以图搜图接口设置频率限制。",
    ):
        add_bullet(doc, text)
    add_heading(doc, "8.2 文件与隐私", 2)
    add_table(
        doc,
        ["对象", "控制措施"],
        [
            ["营业执照", "PRIVATE；仅商家本人和有审核权限的管理员访问；日志中不记录文件内容。"],
            ["地址与手机号", "接口最小返回，列表页面脱敏；导出和日志禁止包含完整值。"],
            ["商品图片", "PUBLIC_READ 或受控 CDN；上传时重编码并去除 EXIF 定位信息。"],
            ["评价图", "审核后公开；限制类型、大小、数量，并执行安全扫描。"],
            ["发货凭证", "仅订单双方与管理员可访问；使用短时签名 URL。"],
        ],
        [1.35, 5.75],
        {0},
    )
    add_heading(doc, "8.3 审计与日志", 2)
    add_body(doc, "operation_logs 保存关键业务操作，不保存密码、令牌、验证码、完整证件号或私有文件 URL。应用日志采用结构化字段，包括 timestamp、level、request_id、user_id、route、latency_ms 和 error_code。审核和订单状态变化需同时具备业务记录与操作日志，两者通过对象编号和 request_id 关联。")
    add_heading(doc, "8.4 常见风险与控制", 2)
    add_table(
        doc,
        ["风险", "控制"],
        [
            ["越权访问", "服务端资源归属校验；集成测试覆盖跨用户、跨商家访问。"],
            ["价格篡改", "后端读取当前报价并重新计算金额，订单保存成交快照。"],
            ["恶意上传", "白名单 MIME、魔数校验、尺寸限制、重编码、随机存储键。"],
            ["SQL 注入", "使用 ORM/参数化查询；禁止字符串拼接 SQL。"],
            ["敏感信息泄漏", "脱敏响应、最小日志、私有存储、配置与密钥不入库。"],
        ],
        [1.45, 5.65],
        {0},
    )

    add_major_break(doc)
    add_heading(doc, "9 部署与运行配置", 1)
    add_heading(doc, "9.1 第一版部署拓扑", 2)
    add_code(doc, [
        "浏览器 / 微信小程序",
        "        ↓ HTTPS",
        "Nginx（静态管理端 + /api 反向代理）",
        "        ↓",
        "FastAPI 应用（含 retrieval 适配器）",
        "   ↙          ↓          ↘",
        "MySQL       Redis       MinIO / 文件存储",
        "                    + FAISS 索引文件（只读加载）",
    ])
    add_heading(doc, "9.2 环境划分", 2)
    add_table(
        doc,
        ["环境", "用途", "数据要求"],
        [
            ["local", "个人开发与单元测试。", "可使用本地 MySQL/容器和种子数据，不含真实个人信息。"],
            ["test", "接口集成、三端联调、回归与算法评估。", "固定测试账号与可重复数据集。"],
            ["demo", "课程演示与答辩。", "演示前备份，禁止临时手改状态；准备离线可运行方案。"],
        ],
        [1.0, 2.3, 3.8],
        {0},
    )
    add_heading(doc, "9.3 配置项", 2)
    add_table(
        doc,
        ["类别", "示例配置", "管理要求"],
        [
            ["数据库", "DATABASE_URL", "通过环境变量或本地 .env 注入，.env 不提交。"],
            ["令牌", "JWT_SECRET / ACCESS_TOKEN_TTL", "不同环境独立密钥，演示环境也不使用默认弱密钥。"],
            ["文件", "STORAGE_ENDPOINT / BUCKET / MAX_UPLOAD_MB", "账号最小权限，私有桶默认拒绝匿名访问。"],
            ["检索", "MODEL_PATH / INDEX_PATH / INDEX_VERSION / TOP_N", "模型、索引与元数据版本必须匹配。"],
            ["日志", "LOG_LEVEL / LOG_DIR", "默认不输出敏感字段；按大小或日期轮转。"],
        ],
        [1.05, 3.05, 3.0],
        {0},
    )
    add_heading(doc, "9.4 启动与健康检查", 2)
    for text in (
        "启动顺序：MySQL/Redis/存储 → 数据库迁移 → 种子数据 → 加载图像模型与索引 → API → 前端。",
        "GET /health/live 只检查进程；GET /health/ready 检查数据库和索引是否可用。",
        "应用启动时比较 INDEX_VERSION、模型版本与 image_embeddings 元数据，不一致时拒绝把检索标记为就绪。",
        "演示环境每日备份数据库；发布前记录 Git 提交、迁移版本、接口版本和索引版本。",
    ):
        add_bullet(doc, text)

    add_major_break(doc)
    add_heading(doc, "10 测试与质量保证", 1)
    add_heading(doc, "10.1 测试分层", 2)
    add_table(
        doc,
        ["层次", "重点", "示例"],
        [
            ["单元测试", "领域规则与纯函数。", "订单金额、状态机、报价有效期、权限判断、图片预处理。"],
            ["数据库测试", "约束、事务、迁移。", "唯一约束、外键、回滚、库存并发、软删除过滤。"],
            ["接口集成", "鉴权、校验、错误码和数据归属。", "跨商家访问返回 403；重复下单返回同一订单。"],
            ["端到端", "三端关键用户旅程。", "商家认证→商品报价→审核→客户检索下单→发货→评价。"],
            ["算法评估", "数据切分、召回、排序和延迟。", "Recall@1/5/10、mAP@10、P95、覆盖率。"],
            ["安全测试", "越权、上传、输入和敏感信息。", "IDOR、超大文件、伪造 MIME、日志泄漏检查。"],
        ],
        [1.15, 2.25, 3.7],
        {0},
    )
    add_heading(doc, "10.2 第一版必测场景", 2)
    add_table(
        doc,
        ["编号", "场景", "预期"],
        [
            ["T-01", "未认证商家提交商品", "403，数据库无新增商品。"],
            ["T-02", "管理员驳回报价但未填原因", "422，状态不变。"],
            ["T-03", "两个请求同时购买最后一件库存", "最多一个成功，不出现负库存。"],
            ["T-04", "相同幂等键重复创建订单", "返回同一订单，不重复扣库存。"],
            ["T-05", "客户访问他人订单", "404 或 403，不泄漏订单内容。"],
            ["T-06", "上传非图片伪装为 JPG", "拒绝上传，不生成业务关联。"],
            ["T-07", "检索结果包含已下架 SKU", "业务过滤后不返回该项。"],
            ["T-08", "完成订单后重复评价", "第二次返回 409。"],
        ],
        [0.75, 3.35, 3.0],
        {0},
    )
    add_heading(doc, "10.3 合并与发布门槛", 2)
    for text in (
        "主分支合并前：格式检查、静态检查、单元测试和核心接口集成测试通过。",
        "数据库变更必须包含迁移脚本、回滚说明和对现有数据的影响说明。",
        "接口变更必须同步 OpenAPI、前端类型和测试；破坏性变更需提升版本或提供兼容期。",
        "算法版本发布必须附评估报告、模型 checksum、索引版本和失败回退方案。",
        "答辩前执行全流程回归，导出测试结果并保存可恢复的数据库备份。",
    ):
        add_bullet(doc, text)

    add_major_break(doc)
    add_heading(doc, "11 需求追踪与迭代计划", 1)
    add_heading(doc, "11.1 需求到实现映射", 2)
    add_table(
        doc,
        ["需求域", "主要模块", "主要表", "核心测试"],
        [
            ["商家认证", "identity / merchant", "merchant_profiles / merchant_audits", "T-01、审核权限与驳回原因。"],
            ["商品与报价", "catalog / pricing", "products / product_skus / merchant_offers", "状态过滤、唯一有效报价、版本冲突。"],
            ["以图搜图", "search / retrieval", "product_images / image_embeddings", "Recall@K、P95、下架过滤。"],
            ["订单履约", "order / notification", "orders / order_items / shipment_records", "T-03、T-04、T-05。"],
            ["评价", "review / files", "reviews / review_images", "T-06、T-08。"],
            ["管理审计", "audit / admin", "operation_logs / 各审核表", "越权、不可抵赖、日志脱敏。"],
        ],
        [1.2, 1.65, 2.75, 1.5],
        {0},
    )
    add_heading(doc, "11.2 建议开发顺序", 2)
    add_table(
        doc,
        ["阶段", "交付内容", "完成判据"],
        [
            ["M1 基础骨架", "仓库结构、配置、数据库迁移、认证与 RBAC、文件接口。", "三端可登录；角色与数据隔离测试通过。"],
            ["M2 商品报价", "商家认证、分类、SPU/SKU、图片、商品/报价审核。", "客户端能看到审核通过的灯具和有效报价。"],
            ["M3 交易闭环", "下单、快照、库存、商家确认、发货、完成与评价。", "端到端交易流程连续跑通。"],
            ["M4 图像检索", "数据整理、基线模型、索引、接口、结果聚合与实验报告。", "检索可用且指标可复现。"],
            ["M5 加固答辩", "权限/并发/上传测试、界面整理、部署脚本、演示数据和文档。", "全流程回归通过，可离线演示和恢复。"],
        ],
        [1.15, 3.75, 2.2],
        {0},
    )
    add_heading(doc, "11.3 待团队确认事项", 2)
    for text in (
        "客户端和商家端是两个独立小程序，还是同一小程序按角色切换入口。",
        "订单第一版是否需要购物车；当前设计允许从单个报价直接下单。",
        "库存采用确认订单时扣减还是创建订单时预占；本设计建议创建时预占并设置超时释放。",
        "课程是否指定数据库、前端框架或部署环境；如有硬性要求需替换对应技术但保留领域边界。",
        "图像数据来源、可用数量、标注方式和版权边界；模型选择应在数据盘点后冻结。",
    ):
        add_bullet(doc, text)
    add_note(doc, "下一步", "以本设计、schema.sql 和 openapi-v0.1.yaml 为基线创建代码骨架，先完成 M1 的认证、RBAC、数据库迁移和健康检查，再进入商品与报价模块。")

    doc.core_properties.title = "灯具商城系统设计说明书"
    doc.core_properties.subject = "三端灯具商城与以图搜图系统设计"
    doc.core_properties.author = "灯具商城项目组"
    doc.core_properties.keywords = "灯具商城, 系统设计, FastAPI, Vue, uni-app, 图像检索"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
