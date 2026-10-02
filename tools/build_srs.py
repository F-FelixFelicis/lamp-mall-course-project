from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\软件工程\灯具商城项目")
OUTPUT = ROOT / "文档" / "灯具商城需求规格说明书_v0.1.docx"

NAVY = "203864"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"
TEXT_GRAY = RGBColor(89, 89, 89)


def set_run_font(run, east_asia="宋体", latin="Arial", size=10.5, bold=None, color=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=LIGHT_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def format_cell(cell, header=False, center=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_borders(cell)
    set_cell_margins(cell)
    if header:
        shade_cell(cell, NAVY)
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center or header else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.12
        if header:
            p.paragraph_format.keep_with_next = True
        for run in p.runs:
            set_run_font(
                run,
                east_asia="微软雅黑" if header else "宋体",
                size=9.2,
                bold=header,
                color=RGBColor(255, 255, 255) if header else RGBColor(0, 0, 0),
            )


def add_table(doc, headers, rows, widths=None, center_cols=None):
    center_cols = set(center_cols or [])
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, value in enumerate(headers):
        hdr.cells[i].text = str(value)
        if widths:
            set_cell_width(hdr.cells[i], widths[i])
        format_cell(hdr.cells[i], header=True, center=True)
    for r_idx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
            if widths:
                set_cell_width(cells[i], widths[i])
            if r_idx % 2 == 1:
                shade_cell(cells[i], PALE_BLUE)
            format_cell(cells[i], center=i in center_cols)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    p.paragraph_format.keep_with_next = True
    return p


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead, rest = text[: len(bold_lead)], text[len(bold_lead) :]
        r = p.add_run(lead)
        set_run_font(r, east_asia="宋体", size=10.5, bold=True)
        r = p.add_run(rest)
        set_run_font(r, east_asia="宋体", size=10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, east_asia="宋体", size=10.5)
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(6)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    r = p.add_run(text)
    set_run_font(r, east_asia="宋体", size=10.5)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(3)
    return p


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, east_asia="宋体", size=9, color=TEXT_GRAY)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    if title._element.pPr is not None:
        title_border = title._element.pPr.find(qn("w:pBdr"))
        if title_border is not None:
            title._element.pPr.remove(title_border)

    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(10.5)


def configure_sections(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.72)
        section.bottom_margin = Inches(0.88)
        section.left_margin = Inches(0.82)
        section.right_margin = Inches(0.82)
        section.footer_distance = Inches(0.32)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("灯具商城需求规格说明书 v0.1    第 ")
    set_run_font(r, east_asia="宋体", size=9, color=TEXT_GRAY)
    add_page_field(p)
    r = p.add_run(" 页")
    set_run_font(r, east_asia="宋体", size=9, color=TEXT_GRAY)


def page_break(doc):
    # Major sections normally flow naturally. Explicit break paragraphs can be
    # pushed onto an otherwise empty page when a preceding table fills a page.
    return None


def build():
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    add_footer(doc.sections[0])

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(82)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("灯具商城需求规格说明书")
    set_run_font(r, east_asia="微软雅黑", size=27, bold=True)
    p.style = doc.styles["Title"]
    remove_paragraph_borders(p)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run("课程综合类项目第一版范围")
    set_run_font(r, east_asia="微软雅黑", size=15, bold=False, color=TEXT_GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(90)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in ("文档版本  v0.1", "编制日期  2026年9月24日", "项目形态  客户端小程序  商家端小程序  PC管理端"):
        r = p.add_run(line + "\n")
        set_run_font(r, east_asia="宋体", size=11)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(66)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("本版用于确认课程范围  业务闭环  算法任务和验收标准")
    set_run_font(r, east_asia="宋体", size=10.5, color=TEXT_GRAY)

    doc.add_page_break()
    add_heading(doc, "文档控制", 1)
    add_table(
        doc,
        ["版本", "日期", "状态", "说明"],
        [["v0.1", "2026-09-24", "范围基线", "根据课程要求和灯具商城附件整理第一版需求，统一家具示例为灯具领域。"]],
        [0.75, 1.15, 1.0, 4.2],
        {0, 1, 2},
    )
    add_heading(doc, "编制依据", 2)
    add_bullet(doc, "课程要求.docx：综合类项目、过程材料、最终提交、测试和现场答辩要求。")
    add_bullet(doc, "附件二：灯具商城.pdf：三端角色、商品报价、订单、采购、发货和评价等业务需求。")
    add_bullet(doc, "附件中的家具名称和分类仅作为原型示例，本说明书统一替换为灯具领域概念。")
    add_heading(doc, "目录", 2)
    for item in (
        "1 项目目标与范围", "2 用户角色与权限", "3 核心业务流程", "4 功能需求",
        "5 数据需求", "6 以图搜图算法需求", "7 外部接口需求", "8 非功能需求",
        "9 验收标准", "10 版本范围与后续扩展", "11 需求追踪与变更控制",
    ):
        add_bullet(doc, item)

    page_break(doc)
    add_heading(doc, "1 项目目标与范围", 1)
    add_heading(doc, "1.1 项目定位", 2)
    add_body(doc, "本项目建设一个面向灯具零售、批量采购和定制询价场景的多商家商城。系统以灯具商品和规格为基础，以商家报价为交易依据，为客户提供关键词搜索、以图搜图、报价比较、下单、物流查看和评价功能，并为商家和管理员提供入驻审核、商品发布、报价审核、订单处理和运营管理能力。")
    add_heading(doc, "1.2 第一版目标", 2)
    for text in (
        "建立客户、商家、管理员三类角色及相互隔离的权限体系。",
        "建立灯具 SPU、SKU、商家报价、库存、订单、发货凭证和评价的数据链路。",
        "实现可复现实验的以图搜图功能，并将检索结果直接关联可购买或可询价的 SKU。",
        "跑通商家入驻、商品发布、报价审核、客户检索、下单、发货和评价的完整流程。",
        "保留 Git 提交、测试结果、算法实验和需求变更记录，满足课程答辩检查要求。",
    ):
        add_bullet(doc, text)
    add_heading(doc, "1.3 产品边界", 2)
    add_table(
        doc,
        ["范围类型", "第一版内容"],
        [
            ["纳入", "账号与角色、商家认证、灯具分类、SPU/SKU、报价与审核、文字搜索、以图搜图、订单、发货凭证、评价、管理后台。"],
            ["暂缓", "在线支付、退款售后、实时聊天、直播、优惠券、复杂营销、跨境结算、完整仓储配送系统、视频评价。"],
            ["不采用", "将第三方识图接口直接作为唯一算法成果；将家具分类和字段原样复制到灯具商城。"],
        ],
        [1.05, 6.05],
        {0},
    )
    add_heading(doc, "1.4 成功标准", 2)
    add_body(doc, "第一版成功的判定标准是核心流程可以连续演示，数据状态前后一致，以图搜图能够返回可解释的 Top-K 结果，测试和实验结果可以重复，团队成员能够在答辩现场说明并修改核心代码。")

    page_break(doc)
    add_heading(doc, "2 用户角色与权限", 1)
    add_table(
        doc,
        ["角色", "主要目标", "核心权限", "禁止操作"],
        [
            ["客户", "发现并购买或询价灯具", "浏览与检索、收藏、查看报价、创建订单、查看物流、提交评价、维护收货地址", "不得发布商品、修改报价或访问后台数据"],
            ["商家", "发布灯具并处理报价与订单", "提交认证、维护本店商品和 SKU、报价、上下架、处理订单、上传发货凭证", "不得直接通过自身商品审核、修改其他商家数据"],
            ["管理员", "维护平台秩序和交易数据", "用户与权限、商家审核、商品审核、报价审核、价格规则、订单、评论、轮播和客服信息管理", "不得绕过审计直接篡改历史操作记录"],
        ],
        [0.75, 1.45, 3.35, 1.55],
        {0},
    )
    add_heading(doc, "2.1 身份与权限原则", 2)
    for text in (
        "系统使用基于角色的访问控制。接口端和页面端均需校验权限，不能只隐藏按钮。",
        "商家认证、商品审核和报价审核必须记录操作者、处理时间、处理结果和驳回原因。",
        "商家只能读写自身店铺、商品、报价和订单范围内的数据。",
        "客户身份证明、手机号、地址和营业执照等敏感字段按最小必要原则展示。",
    ):
        add_bullet(doc, text)

    page_break(doc)
    add_heading(doc, "3 核心业务流程", 1)
    add_heading(doc, "3.1 商家入驻与商品发布", 2)
    add_body(doc, "商家注册后提交法人姓名、营业执照、联系电话、地址和店铺名称。管理员审核通过后，商家可以创建灯具商品，维护 SKU 图片、规格、库存和报价。商品或关键字段发生变化时重新进入审核，审核通过后才对客户可见。")
    add_heading(doc, "3.2 多商家报价与价格展示", 2)
    add_body(doc, "同一 SKU 可以有多个商家报价。平台保存每条报价的商家、价格、库存、状态、生效时间和备注。客户侧默认展示当前有效最低报价，并可以展开查看全部有效报价。没有有效商家报价时，可展示管理员维护的平台参考价，但必须明确标识为参考价。")
    add_heading(doc, "3.3 搜索与以图搜图", 2)
    add_body(doc, "客户可以按商品名称、SKU、分类和属性搜索，也可以上传灯具图片。图像检索服务返回相似 SKU 及相似度，业务服务再过滤未上架、审核未通过或无有效报价的数据。客户从结果进入商品详情后，可以查看原图、规格、库存和商家报价。")
    add_heading(doc, "3.4 下单与履约", 2)
    add_body(doc, "客户选择 SKU、商家报价、数量和收货方式后创建订单。商家确认订单并更新备货与发货状态，发货时上传物流单或提货凭证。客户可以查看状态和凭证，在订单完成后提交文字或图片评价。")
    add_heading(doc, "3.5 核心状态", 2)
    add_table(
        doc,
        ["对象", "状态顺序", "约束"],
        [
            ["商家认证", "待审核 → 已通过 或 已驳回", "驳回必须填写原因；重新提交生成新的审核记录。"],
            ["商品", "草稿 → 待审核 → 已上架 → 已下架", "待审核和已驳回商品不得在客户端展示。"],
            ["报价", "草稿 → 待审核 → 已生效 → 已失效", "同一商家同一 SKU 同时只允许一条生效报价。"],
            ["订单", "待确认 → 待发货 → 已发货 → 已完成 或 已取消", "状态只能按允许的路径变化，并记录操作日志。"],
        ],
        [1.0, 2.55, 3.55],
        {0},
    )

    page_break(doc)
    add_heading(doc, "4 功能需求", 1)
    add_heading(doc, "4.1 通用账号与基础能力", 2)
    common_rows = [
        ["FR-COM-001", "注册登录", "支持手机号与验证码注册；支持账号密码登录。微信授权登录作为可选扩展。", "有效用户可以进入对应角色端；错误凭证给出明确提示。"],
        ["FR-COM-002", "密码重置", "用户通过手机号验证码设置新密码。", "重置后旧密码失效；验证码过期后不可使用。"],
        ["FR-COM-003", "个人资料", "用户维护昵称、姓名、手机号和头像；客户维护收货地址。", "字段校验通过后保存；敏感字段按权限脱敏显示。"],
        ["FR-COM-004", "文件上传", "支持商品图、营业执照、评价图和发货凭证上传。", "校验类型、大小和数量；失败不产生无效业务记录。"],
        ["FR-COM-005", "消息提醒", "对审核结果、订单状态、发货和未读评价提供站内提醒。", "提醒可标记已读，并能跳转到对应业务对象。"],
    ]
    add_table(doc, ["编号", "名称", "需求说明", "验收要点"], common_rows, [1.05, 1.0, 3.15, 1.9], {0, 1})

    add_heading(doc, "4.2 客户端小程序", 2)
    customer_rows = [
        ["FR-CUS-001", "首页与分类", "展示轮播、灯具分类和推荐商品，支持进入分类列表。", "只展示已上架且审核通过的商品。"],
        ["FR-CUS-002", "文字搜索", "按名称、SKU、分类、风格、材质、功率和色温等条件检索。", "支持分页、排序和无结果提示。"],
        ["FR-CUS-003", "以图搜图", "上传图片并获得相似灯具 Top-K 结果。", "结果包含商品、SKU、相似度和最低有效报价。"],
        ["FR-CUS-004", "商品详情", "展示商品图片、SKU 属性、库存、参考价、最低报价及商家报价列表。", "切换 SKU 后图片、库存和报价同步变化。"],
        ["FR-CUS-005", "收藏", "客户收藏或取消收藏商品。", "收藏列表与商品状态保持一致。"],
        ["FR-CUS-006", "询价与联系", "批量或定制需求可进入询价说明并查看平台客服信息。", "记录目标商品和 SKU，避免客服无法定位需求。"],
        ["FR-CUS-007", "创建订单", "选择 SKU、商家、数量、收货方式和地址后创建订单。", "后端重新校验商品、报价和库存；保存价格快照。"],
        ["FR-CUS-008", "订单查询", "查看订单列表、详情、状态和发货凭证。", "客户只能查看本人订单。"],
        ["FR-CUS-009", "商品评价", "完成订单后提交文字和图片评价。", "每个订单明细只能形成一条有效评价。"],
    ]
    add_table(doc, ["编号", "名称", "需求说明", "验收要点"], customer_rows, [1.05, 1.0, 3.15, 1.9], {0, 1})

    page_break(doc)
    add_heading(doc, "4.3 商家端小程序", 2)
    merchant_rows = [
        ["FR-MER-001", "商家认证", "提交法人、营业执照、电话、地址和店铺名称。", "未通过认证不能发布商品或报价。"],
        ["FR-MER-002", "商品维护", "创建和编辑本店商品，维护灯具图片、描述和分类。", "关键字段变化后重新审核。"],
        ["FR-MER-003", "SKU 维护", "维护颜色、尺寸、材质、功率、色温、光通量、库存和图片。", "SKU 编码唯一；库存不能为负。"],
        ["FR-MER-004", "商家报价", "针对平台 SKU 提交、修改或撤回报价，并填写库存与备注。", "修改后的报价重新审核；保留历史版本。"],
        ["FR-MER-005", "上下架", "控制已审核商品是否对客户可见。", "下架不删除历史订单和报价快照。"],
        ["FR-MER-006", "订单处理", "查询本店订单，确认订单并更新备货、发货状态。", "状态变化符合订单状态机。"],
        ["FR-MER-007", "发货凭证", "上传物流单图片或填写物流信息。", "上传后客户可在订单详情查看。"],
        ["FR-MER-008", "审核反馈", "查看商品和报价驳回原因、未读提醒并重新提交。", "重新提交产生新的审核记录。"],
    ]
    add_table(doc, ["编号", "名称", "需求说明", "验收要点"], merchant_rows, [1.05, 1.0, 3.15, 1.9], {0, 1})

    admin_heading = add_heading(doc, "4.4 PC 管理端", 2)
    admin_heading.paragraph_format.page_break_before = True
    admin_rows = [
        ["FR-ADM-001", "用户与权限", "管理后台账号、角色、权限和启停状态。", "权限变更即时生效并记录审计日志。"],
        ["FR-ADM-002", "商家审核", "查看商家认证材料，执行通过或驳回。", "驳回必须填写原因。"],
        ["FR-ADM-003", "商品审核", "审核商家提交的商品及 SKU，支持查询、上下架和软删除。", "审核记录可追踪。"],
        ["FR-ADM-004", "报价审核", "审核商家报价，查看同 SKU 全部报价和历史变化。", "审核通过后才参与最低价计算。"],
        ["FR-ADM-005", "价格规则", "维护平台参考价、利润率或展示规则。", "规则变化记录版本，不回写历史订单。"],
        ["FR-ADM-006", "客户管理", "按姓名、手机号或订单查询客户及必要信息。", "地址等敏感信息默认脱敏。"],
        ["FR-ADM-007", "订单管理", "查询全部订单、查看状态、处理异常并导出课程演示数据。", "关键修改需要操作原因和审计记录。"],
        ["FR-ADM-008", "评价管理", "查看、回复、隐藏违规评价并处理未读提醒。", "隐藏操作不物理删除原始记录。"],
        ["FR-ADM-009", "运营配置", "维护轮播图、链接、客服二维码和灯具分类。", "配置发布后客户端刷新可见。"],
        ["FR-ADM-010", "统计概览", "展示商家、商品、报价、订单及以图搜图调用量。", "统计口径与明细数据一致。"],
    ]
    add_table(doc, ["编号", "名称", "需求说明", "验收要点"], admin_rows, [1.05, 1.0, 3.15, 1.9], {0, 1})

    page_break(doc)
    add_heading(doc, "5 数据需求", 1)
    add_heading(doc, "5.1 核心数据对象", 2)
    data_rows = [
        ["用户", "账号、手机号、密码摘要、角色、状态", "手机号唯一；密码不得明文保存。"],
        ["商家", "店铺名称、法人、营业执照、联系方式、地址、认证状态", "一个用户绑定一个主要商家主体。"],
        ["商品 SPU", "名称、分类、品牌、风格、描述、状态", "描述一个可销售的灯具款式。"],
        ["商品 SKU", "SKU 编码、颜色、尺寸、材质、功率、色温、库存、图片", "编码唯一；属于一个 SPU。"],
        ["商家报价", "商家、SKU、价格、库存、状态、生效时间、备注", "同商家同 SKU 同时仅一条生效记录。"],
        ["图像向量", "SKU 图片、模型版本、向量、生成时间", "模型版本变化后可重建索引。"],
        ["订单", "客户、商家、金额、地址快照、状态、创建时间", "金额和地址使用下单时快照。"],
        ["订单明细", "SKU、报价快照、数量、小计", "不能依赖后续变化的商品价格。"],
        ["发货记录", "订单、物流信息、凭证图片、发货时间", "与订单状态变化保持一致。"],
        ["评价", "订单明细、评分、文字、图片、状态", "仅完成订单可评价。"],
        ["审核记录", "对象类型、对象编号、结果、原因、操作者、时间", "审核历史不可被业务修改覆盖。"],
        ["操作日志", "用户、动作、对象、时间、结果、必要上下文", "用于答辩演示和问题追踪。"],
    ]
    add_table(doc, ["对象", "主要字段", "关键约束"], data_rows, [1.15, 3.45, 2.55], {0})
    add_heading(doc, "5.2 灯具属性字典", 2)
    add_body(doc, "第一版至少支持分类、品牌、应用空间、风格、材质、颜色、尺寸、功率、色温、光通量、电压、光源类型、调光方式、安装方式和智能协议。分类与可选属性应由后台维护，避免将家具示例字段或固定枚举写死在客户端。")

    page_break(doc)
    add_heading(doc, "6 以图搜图算法需求", 1)
    add_heading(doc, "6.1 问题定义", 2)
    add_body(doc, "输入是一张用户上传的灯具图片，输出是按视觉相似度排序的 Top-K 商品 SKU。系统需要处理背景干扰、拍摄角度、光照、裁剪和同款不同颜色等差异，并在业务层过滤未审核、未上架和无有效报价的数据。")
    add_heading(doc, "6.2 算法流程", 2)
    for text in (
        "离线处理商品主图，完成尺寸归一化、特征提取、向量归一化和索引构建。",
        "在线接收用户图片，完成格式校验、预处理和查询向量生成。",
        "在向量索引中召回候选 SKU，使用余弦相似度或等价距离排序。",
        "结合商品状态、分类、库存和报价完成业务过滤，返回 Top-K 结果。",
        "记录模型版本、查询耗时和匿名化结果，用于实验复现和错误分析。",
    ):
        add_bullet(doc, text)
    add_heading(doc, "6.3 基线与实验", 2)
    add_table(
        doc,
        ["实验项", "第一版要求"],
        [
            ["数据集", "建立课程用灯具图片集，按 SKU 或商品款式标注相关结果，并划分索引集与查询集。"],
            ["基线方法", "实现颜色直方图或颜色与纹理组合特征，作为传统方法基线。"],
            ["主方法", "使用可解释的深度图像特征提取模型生成向量，记录模型名称、版本、输入尺寸和向量维度。"],
            ["评价指标", "至少报告 Precision@K、Recall@K、平均查询耗时和 P95 查询耗时。"],
            ["对比分析", "比较基线与主方法，分析灯具类别、背景、角度和颜色变化造成的成功与失败案例。"],
            ["复现要求", "固定数据划分与随机种子；提供索引构建、评测脚本和实验配置。"],
        ],
        [1.15, 6.0],
        {0},
    )
    add_heading(doc, "6.4 建议验收目标", 2)
    add_table(
        doc,
        ["指标", "v0.1 建议目标", "说明"],
        [
            ["Recall@5", "不低于 0.75", "相关商品在前 5 个结果中至少命中一次的查询比例。"],
            ["Precision@5", "不低于 0.60", "前 5 个结果中相关商品所占比例。"],
            ["在线查询耗时", "P95 不高于 2 秒", "不含用户网络上传时间，使用课程演示环境测量。"],
            ["可用性", "无匹配时返回空结果及引导", "不得用不相关商品填充结果。"],
        ],
        [1.35, 1.55, 4.25],
        {0, 1},
    )
    add_body(doc, "上述数值是第一版建议目标。团队完成数据集初测后，可以通过需求变更记录调整阈值，但不得取消基线比较、实验复现和失败案例分析。")

    page_break(doc)
    add_heading(doc, "7 外部接口需求", 1)
    add_table(
        doc,
        ["接口类别", "用途", "第一版约束"],
        [
            ["短信验证码", "注册和密码重置", "开发环境允许使用可审计的模拟验证码；生产配置不得写入仓库。"],
            ["对象存储", "商品图、执照、评价图和发货凭证", "使用受控文件类型与访问地址；敏感文件不公开。"],
            ["图像检索服务", "特征提取和 Top-K 召回", "以自有评测流程为准；第三方服务只能作为可替换组件或对照。"],
            ["微信授权", "小程序快捷登录", "列为可选扩展，不阻塞账号密码主流程。"],
        ],
        [1.35, 2.15, 3.65],
        {0},
    )
    add_heading(doc, "7.1 接口通用约定", 2)
    for text in (
        "接口使用统一的身份认证、错误码、分页参数和时间格式。",
        "写操作需要校验请求参数、角色权限和业务状态，重复请求不得产生重复订单或重复审核。",
        "文件上传先获得受控上传凭证，再提交业务对象，避免孤立文件无限增长。",
        "接口文档应包含请求、响应、权限、错误场景和示例，并与实现版本保持一致。",
    ):
        add_bullet(doc, text)

    page_break(doc)
    add_heading(doc, "8 非功能需求", 1)
    nfr_rows = [
        ["NFR-001", "性能", "普通列表和详情接口在课程演示环境下 P95 不高于 800 毫秒；图像检索按第 6.4 节执行。"],
        ["NFR-002", "安全", "密码使用强哈希；接口鉴权；上传白名单；防止越权、SQL 注入和恶意文件上传。"],
        ["NFR-003", "隐私", "客户地址、手机号、身份证明和营业执照仅向必要角色展示，日志不得记录完整敏感值。"],
        ["NFR-004", "一致性", "创建订单、扣减或锁定库存、保存价格快照应在一致的事务边界内完成。"],
        ["NFR-005", "可用性", "关键失败操作给出可理解提示；外部服务不可用时不破坏已有订单和商品数据。"],
        ["NFR-006", "可维护性", "按账号、商品、报价、订单、检索等领域划分模块，禁止将核心规则只写在页面组件中。"],
        ["NFR-007", "可测试性", "核心服务可独立测试；测试数据可重建；图像检索实验具有固定配置。"],
        ["NFR-008", "审计", "商家审核、商品审核、报价审核、订单异常修改和权限变更保留操作记录。"],
        ["NFR-009", "兼容性", "客户端适配主流微信小程序环境；PC 管理端适配课程演示使用的现代浏览器。"],
    ]
    add_table(doc, ["编号", "类别", "要求"], nfr_rows, [1.05, 1.0, 5.1], {0, 1})

    page_break(doc)
    add_heading(doc, "9 验收标准", 1)
    acceptance_rows = [
        ["AC-01", "商家入驻", "商家提交认证，管理员驳回并填写原因，商家修改后再次提交并通过。"],
        ["AC-02", "商品与报价", "商家创建灯具和 SKU，提交报价，管理员审核后客户端显示最低价和全部有效报价。"],
        ["AC-03", "以图搜图", "上传测试灯具图片，返回 Top-K SKU；结果可进入详情并查看有效报价；实验指标可复现。"],
        ["AC-04", "订单履约", "客户下单，商家确认并发货，客户查看凭证并完成订单。"],
        ["AC-05", "评价", "完成订单后客户提交评价，管理员处理违规评价且保留原始记录。"],
        ["AC-06", "权限", "客户不能访问商家或管理员接口，商家不能操作其他商家数据。"],
        ["AC-07", "数据一致性", "下架或改价后历史订单仍显示下单时的商品、报价和地址快照。"],
        ["AC-08", "工程材料", "Git 历史、需求、设计、数据库与接口、测试、算法实验、总结和答辩材料齐全。"],
    ]
    add_table(doc, ["编号", "场景", "通过条件"], acceptance_rows, [1.05, 1.35, 4.75], {0, 1})

    add_heading(doc, "9.1 测试范围", 2)
    for text in (
        "单元测试覆盖价格选择、权限判断、状态转换和检索结果过滤等核心规则。",
        "接口测试覆盖注册登录、商品、报价、审核、订单、上传和评价接口。",
        "系统测试按 AC-01 至 AC-08 执行完整场景并保存结果。",
        "性能测试记录接口 P95、图像检索延迟、测试环境、样本量和失败率。",
        "缺陷记录包含发现版本、严重程度、复现步骤、修复提交和回归结果。",
    ):
        add_bullet(doc, text)

    page_break(doc)
    add_heading(doc, "10 版本范围与后续扩展", 1)
    add_table(
        doc,
        ["阶段", "交付范围", "进入下一阶段条件"],
        [
            ["M1 需求基线", "角色、流程、数据对象、需求编号、算法指标和验收标准", "小组确认范围并建立 Git 仓库与任务分工。"],
            ["M2 基础工程", "账号权限、商家认证、分类、SPU/SKU、文件上传", "接口测试通过，权限边界可验证。"],
            ["M3 交易闭环", "报价审核、搜索、详情、订单、发货和评价", "AC-01、AC-02、AC-04、AC-05 通过。"],
            ["M4 算法集成", "数据集、基线、深度特征、向量检索、业务过滤和实验报告", "AC-03 通过并达到确认后的指标。"],
            ["M5 验收交付", "系统测试、性能测试、缺陷修复、文档、演示数据和答辩 PPT", "AC-01 至 AC-08 全部通过。"],
        ],
        [1.25, 3.45, 2.45],
        {0},
    )
    add_heading(doc, "10.1 后续可选扩展", 2)
    add_body(doc, "第一版稳定后再评估在线支付与退款、优惠券、实时聊天、视频评价、推荐系统、复杂采购审批、库存预占、多仓库和多语言。扩展项必须通过需求变更记录进入计划，不能挤占核心闭环和算法实验的完成时间。")

    add_heading(doc, "11 需求追踪与变更控制", 1)
    add_body(doc, "每个功能需求使用唯一编号，并在设计文档、接口文档、代码提交和测试用例中引用。需求发生变化时，记录提出人、原因、影响范围、处理结论和对应版本。涉及角色权限、订单状态、价格计算或算法指标的变更，需要小组共同确认后实施。")
    add_table(
        doc,
        ["需求编号", "设计模块", "接口或服务", "测试用例", "状态"],
        [
            ["示例 FR-CUS-003", "图像检索模块", "POST /image-search", "TC-IMG-001 至 TC-IMG-005", "待建立"],
            ["示例 FR-CUS-007", "订单模块", "POST /orders", "TC-ORD-001 至 TC-ORD-008", "待建立"],
        ],
        [1.25, 1.45, 1.7, 1.75, 0.95],
        {0, 4},
    )
    add_body(doc, "本说明书作为 v0.1 范围基线。数据库设计、接口设计和任务分工应以本说明书的需求编号为入口继续细化。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
